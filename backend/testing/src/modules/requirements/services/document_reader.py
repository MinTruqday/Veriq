import csv
import hashlib
import io
from datetime import datetime, timezone
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import load_workbook
from pymongo.errors import DuplicateKeyError

from src.clients.ai_automation import ai_automation_client
from src.clients.document_conversion import read_requirement_source_image
from src.clients.worker import WorkerClientError, worker_client
from src.core.common import get_project
from src.core.database import database
from src.modules.requirements.services.requirement_documents import read_raw_requirement_source
from src.modules.requirements.services.requirement_documents import (
    can_read_requirement_document,
    get_requirement_document_record,
)
from src.repositories.requirement_document import requirement_document_repository
from src.repositories.requirement_source_block import requirement_source_block_repository


SUMMARY_INPUT_LIMIT = 12000
SUMMARY_PROMPT_VERSION = "1"


class SummaryJobCancelled(Exception):
    pass


def _markdown_table(rows):
    values = [[str(value or "").replace("|", "\\|").replace("\n", " ") for value in row] for row in rows]
    width = max((len(row) for row in values), default=0)
    values = [row + [""] * (width - len(row)) for row in values]
    if not values:
        return ""
    header = values[0]
    body = values[1:]
    return "\n".join(
        [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(["---"] * width) + " |",
            *["| " + " | ".join(row) + " |" for row in body],
        ]
    )


def _search_hits(text, query, limit, offset=0, page_no=None, sheet=None, row_offset=None):
    folded_text = text.casefold()
    folded_query = query.casefold()
    matches = []
    cursor = 0
    total = 0
    while True:
        found = folded_text.find(folded_query, cursor)
        if found < 0:
            break
        total += 1
        if len(matches) < limit:
            start = max(0, found - 100)
            end = min(len(text), found + len(query) + 180)
            matches.append(
                {
                    "offset": offset + found,
                    "page_no": page_no,
                    "sheet": sheet,
                    "row_offset": row_offset,
                    "snippet": text[start:end].strip(),
                }
            )
        cursor = found + len(folded_query)
    return total, matches


async def _document_text(document):
    if document.get("format") == "xlsx":
        data = await read_raw_requirement_source(document)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections = []
        try:
            for worksheet in workbook.worksheets:
                rows = [
                    ",".join(str(value or "") for value in row)
                    for row in worksheet.iter_rows(values_only=True)
                ]
                sections.append(f"Sheet: {worksheet.title}\n" + "\n".join(rows))
        finally:
            workbook.close()
        return "\n\n".join(sections)
    content = document.get("normalized_content") or ""
    if isinstance(content, dict):
        from src.modules.requirements.services.requirement_workflow import serialized_content

        content = serialized_content(content)
    return str(content)


def split_source_text(text: str, limit: int = SUMMARY_INPUT_LIMIT) -> list[str]:
    paragraphs = [part.strip() for part in text.splitlines() if part.strip()]
    chunks = []
    current = ""
    for paragraph in paragraphs:
        while len(paragraph) > limit:
            if current:
                chunks.append(current)
                current = ""
            chunks.append(paragraph[:limit])
            paragraph = paragraph[limit:]
        candidate = f"{current}\n{paragraph}" if current else paragraph
        if len(candidate) > limit:
            chunks.append(current)
            current = paragraph
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


async def read_document_page(document_id, offset, limit, user, sheet=None, row_limit=100):
    document = await get_requirement_document_record(document_id, user)
    if document.get("status") != "READY":
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_CONTENT_NOT_READY"})
    raw = document.get("normalized_content")
    if isinstance(raw, dict):
        from src.modules.requirements.services.requirement_workflow import serialized_content

        raw = serialized_content(raw)
    content = str(raw or "")
    if not content and document.get("format") != "xlsx":
        raise HTTPException(status_code=422, detail={"code": "DOCUMENT_CONTENT_UNAVAILABLE"})
    if document.get("format") == "xlsx":
        data = await read_raw_requirement_source(document)
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            names = workbook.sheetnames
            selected_sheet = sheet if sheet in names else names[0]
            worksheet = workbook[selected_sheet]
            header_rows = list(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
            body_rows = list(
                worksheet.iter_rows(
                    min_row=offset + 2,
                    max_row=min(offset + row_limit + 1, worksheet.max_row or offset + row_limit + 1),
                    values_only=True,
                )
            )
            rows = header_rows[:1] + body_rows
            content = _markdown_table(rows)
            total_rows = worksheet.max_row or 0
            data_row_count = max(total_rows - 1, 0)
            next_offset = offset + len(body_rows) if offset + len(body_rows) < data_row_count else None
            workbook.close()
        except Exception as error:
            raise HTTPException(status_code=422, detail={"code": "DOCUMENT_SHEET_READ_FAILED"}) from error
        return {
            "document_id": str(document["_id"]),
            "project_id": str(document["project_id"]),
            "filename": document.get("filename"),
            "format": document.get("format"),
            "source_version": document.get("source_version") or "1",
            "revision": document.get("revision"),
            "content_hash": document.get("content_hash"),
            "content": content,
            "offset": offset,
            "total_characters": len(content),
            "total_rows": total_rows,
            "next_offset": next_offset,
            "sheet": selected_sheet,
            "sheets": names,
            "page_count": None,
            "unprocessed_blocks": [],
            "extraction_warnings": [],
        }
    if document.get("format") == "csv":
        rows = list(csv.reader(io.StringIO(content)))
        header = rows[0] if rows else []
        body = rows[offset + 1 : offset + row_limit + 1]
        selected_rows = ([header] if header else []) + body
        content = _markdown_table(selected_rows)
        total_rows = len(rows)
        next_offset = offset + len(body) if offset + len(body) < max(total_rows - 1, 0) else None
        report = document.get("parse_report") or {}
        return {
            "document_id": str(document["_id"]),
            "project_id": str(document["project_id"]),
            "filename": document.get("filename"),
            "format": document.get("format"),
            "source_version": document.get("source_version") or "1",
            "revision": document.get("revision"),
            "content_hash": document.get("content_hash"),
            "content": content,
            "offset": offset,
            "total_characters": len(content),
            "total_rows": total_rows,
            "next_offset": next_offset,
            "sheet": None,
            "sheets": [],
            "page_count": report.get("page_count"),
            "unprocessed_blocks": report.get("unprocessed_blocks") or [],
            "extraction_warnings": report.get("extraction_warnings") or [],
        }
    page = content[offset : offset + limit]
    report = document.get("parse_report") or {}
    return {
        "document_id": str(document["_id"]),
        "project_id": str(document["project_id"]),
        "filename": document.get("filename"),
        "format": document.get("format"),
        "source_version": document.get("source_version") or "1",
        "revision": document.get("revision"),
        "content_hash": document.get("content_hash"),
        "content": page,
        "offset": offset,
        "total_characters": len(content),
        "next_offset": offset + len(page) if offset + len(page) < len(content) else None,
        "page_count": report.get("page_count"),
        "source_pages_available": report.get("structure_available") is True,
        "unprocessed_blocks": report.get("unprocessed_blocks") or [],
        "extraction_warnings": report.get("extraction_warnings") or [],
        "sheet": None,
        "sheets": [],
    }


async def read_document_source_page(document_id, page_no, user):
    document = await get_requirement_document_record(document_id, user)
    if document.get("status") != "READY":
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_CONTENT_NOT_READY"})
    if document.get("format") not in {"pdf", "docx"}:
        raise HTTPException(status_code=422, detail={"code": "DOCUMENT_PAGE_VIEW_UNAVAILABLE"})
    source_hash = str(document.get("source_structure_source_hash") or document.get("content_hash") or "")
    if document.get("source_structure_store") == "blocks":
        blocks = await requirement_source_block_repository.list_version(
            document["project_id"], document_id, source_hash
        )
    else:
        blocks = document.get("source_structure") or []
    if not blocks:
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_PAGE_STRUCTURE_UNAVAILABLE"})
    selected = []
    for index, block in enumerate(blocks):
        if not isinstance(block, dict):
            continue
        locations = block.get("locations") or []
        block_page = block.get("page_no") or next(
            (location.get("page_no") for location in locations if location.get("page_no")), None
        )
        if block_page == page_no:
            selected.append(
                {
                    "block_id": block.get("block_id") or str(index),
                    "block_index": index,
                    "type": block.get("type"),
                    "text": str(block.get("text") or ""),
                    "extraction_status": block.get("extraction_status"),
                    "locations": locations,
                    "table_metadata": block.get("table_metadata"),
                    "image_metadata": block.get("image_metadata"),
                }
            )
    report = document.get("parse_report") or {}
    return {
        "document_id": document_id,
        "project_id": document["project_id"],
        "filename": document.get("filename"),
        "format": document.get("format"),
        "source_version": document.get("source_version") or "1",
        "content_hash": document.get("content_hash"),
        "page_no": page_no,
        "page_count": report.get("page_count"),
        "source_pages_available": True,
        "blocks": selected,
        "content": "\n\n".join(block["text"] for block in selected if block["text"]),
        "unprocessed_blocks": report.get("unprocessed_blocks") or [],
    }


async def read_document_source_image(document_id, block_index, user):
    document = await get_requirement_document_record(document_id, user)
    if document.get("status") != "READY":
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_CONTENT_NOT_READY"})
    source_hash = str(document.get("source_structure_source_hash") or document.get("content_hash") or "")
    if document.get("source_structure_store") == "blocks":
        rows = await requirement_source_block_repository.page_version(
            document["project_id"], document_id, source_hash, block_index, 1
        )
        block = rows[0].get("block") if rows and rows[0].get("block_index") == block_index else None
    else:
        blocks = document.get("source_structure") or []
        block = blocks[block_index] if 0 <= block_index < len(blocks) else None
    if block_index < 0 or not isinstance(block, dict):
        raise HTTPException(status_code=404, detail={"code": "DOCUMENT_SOURCE_IMAGE_NOT_FOUND"})
    image = block.get("image_metadata") if isinstance(block, dict) else None
    asset_key = image.get("asset_key") if isinstance(image, dict) else None
    if not asset_key:
        raise HTTPException(status_code=404, detail={"code": "DOCUMENT_SOURCE_IMAGE_NOT_FOUND"})
    try:
        data, content_type = await read_requirement_source_image(
            str(document["project_id"]), document_id, asset_key
        )
    except Exception as error:
        raise HTTPException(status_code=404, detail={"code": "DOCUMENT_SOURCE_IMAGE_UNAVAILABLE"}) from error
    return data, content_type


async def search_document(document_id, query, user, limit=50):
    document = await get_requirement_document_record(document_id, user)
    if document.get("status") != "READY":
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_CONTENT_NOT_READY"})
    query = query.strip()
    if len(query) < 2:
        raise HTTPException(status_code=422, detail={"code": "SEARCH_QUERY_TOO_SHORT"})
    hits = []
    total = 0
    if document.get("format") in {"pdf", "docx"}:
        source_hash = str(document.get("source_structure_source_hash") or document.get("content_hash") or "")
        blocks = (
            await requirement_source_block_repository.list_version(
                document["project_id"], document_id, source_hash
            )
            if document.get("source_structure_store") == "blocks"
            else document.get("source_structure") or []
        )
        if blocks:
            for block_index, block in enumerate(blocks):
                if not isinstance(block, dict):
                    continue
                block_text = str(block.get("text") or "")
                locations = block.get("locations") or []
                page_no = block.get("page_no") or next(
                    (location.get("page_no") for location in locations if location.get("page_no")), None
                )
                remaining = max(limit - len(hits), 0)
                count, found = _search_hits(
                    block_text,
                    query,
                    remaining,
                    page_no=page_no,
                )
                for item in found:
                    item["block_id"] = block.get("block_id") or str(block_index)
                    item["block_index"] = block_index
                total += count
                hits.extend(found)
        else:
            total, hits = _search_hits(await _document_text(document), query, limit)
        pages = (document.get("parse_report") or {}).get("page_count")
    elif document.get("format") == "csv":
        content = await _document_text(document)
        rows = list(csv.reader(io.StringIO(content)))
        for row_index, row in enumerate(rows[1:]):
            row_text = " | ".join(row)
            remaining = max(limit - len(hits), 0)
            count, found = _search_hits(
                row_text,
                query,
                remaining,
                row_offset=row_index,
            )
            total += count
            hits.extend(found)
        pages = None
    elif document.get("format") == "xlsx":
        data = await read_raw_requirement_source(document)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                    row_text = " | ".join(str(value or "") for value in row)
                    remaining = max(limit - len(hits), 0)
                    count, found = _search_hits(
                        row_text,
                        query,
                        remaining,
                        sheet=worksheet.title,
                        row_offset=max(0, row_index - 1),
                    )
                    total += count
                    hits.extend(found)
        finally:
            workbook.close()
        pages = None
    else:
        content = await _document_text(document)
        total, hits = _search_hits(content, query, limit)
        pages = (document.get("parse_report") or {}).get("page_count")
    return {
        "document_id": document_id,
        "source_version": document.get("source_version") or "1",
        "content_hash": document.get("content_hash"),
        "query": query,
        "total_matches": total,
        "matches": hits,
        "truncated": total > len(hits),
        "page_count": pages,
    }


async def _assert_summary_source_current(document, user):
    latest = await get_requirement_document_record(document["_id"], user)
    if (
        latest.get("content_hash") != document.get("content_hash")
        or latest.get("revision") != document.get("revision")
        or (latest.get("source_version") or "1") != (document.get("source_version") or "1")
        or latest.get("ai_enabled") is not True
    ):
        raise RuntimeError("SOURCE_CHANGED_OR_ACCESS_REVOKED")


async def _update_summary_job(job_id, changes):
    if job_id:
        await database.value.document_summary_jobs.update_one(
            {"_id": job_id}, {"$set": {**changes, "updated_at": datetime.now(timezone.utc)}}
        )


async def _lock_summary_sources(project_id, request, user):
    if request.all_readable:
        documents = await requirement_document_repository.list({"project_id": project_id}, None)
        candidates = [
            document
            for document in documents
            if can_read_requirement_document(document, user)
            and document.get("status") != "ARCHIVED"
        ]
    else:
        candidates = []
        for document_id in request.document_ids:
            document = await get_requirement_document_record(document_id, user)
            if str(document.get("project_id")) != str(project_id):
                raise HTTPException(status_code=404, detail={"code": "DOCUMENT_NOT_FOUND"})
            candidates.append(document)
    return [
        {
            "document_id": str(document["_id"]),
            "revision": document.get("revision"),
            "content_hash": document.get("content_hash"),
            "source_version": document.get("source_version") or "1",
            "ai_enabled": document.get("ai_enabled") is True,
        }
        for document in candidates
    ]


async def _enqueue_summary_job(job, user, generation=0):
    request = job["request"]
    worker_request = {
        "event": "document.summary.requested",
        "project_id": job["project_id"],
        "artifact_version_id": job["_id"],
        "model_version": f"document-summary-v{SUMMARY_PROMPT_VERSION}",
        "requester_id": user.id,
        "requester_email": user.email,
        "payload": {
            "summary_job_id": job["_id"],
            "request": request,
            "locked_sources": job.get("locked_sources") or [],
            "dispatch_generation": generation,
        },
    }
    worker_job = await worker_client.enqueue(worker_request)
    await database.value.document_summary_jobs.update_one(
        {"_id": job["_id"], "project_id": job["project_id"]},
        {
            "$set": {
                "worker_job_id": worker_job.get("job_id"),
                "worker_status": worker_job.get("status"),
                "updated_at": datetime.now(timezone.utc),
            }
        },
    )
    return worker_job


async def _check_summary_job(job_id):
    if job_id:
        job = await database.value.document_summary_jobs.find_one(
            {"_id": job_id}, {"cancel_requested": 1}
        )
        if not job or job.get("cancel_requested"):
            raise SummaryJobCancelled()


async def process_document_summary_job(body, payload, user):
    project_id = str(body.get("project_id") or "")
    job_id = str(payload.get("summary_job_id") or "")
    worker_job_id = str(body.get("job_id") or "")
    if not project_id or not job_id or not worker_job_id:
        raise HTTPException(status_code=422, detail={"code": "DOCUMENT_SUMMARY_JOB_CONTEXT_INVALID"})
    job = await database.value.document_summary_jobs.find_one(
        {"_id": job_id, "project_id": project_id, "requester_id": user.id}
    )
    if not job:
        raise HTTPException(status_code=404, detail={"code": "DOCUMENT_SUMMARY_JOB_NOT_FOUND"})
    if job.get("cancel_requested"):
        await _update_summary_job(job_id, {"status": "CANCELLED", "completed_at": datetime.now(timezone.utc)})
        return {"status": "COMPLETED", "summary_job_id": job_id, "summary_status": "CANCELLED"}
    claimed = await database.value.document_summary_jobs.update_one(
        {
            "_id": job_id,
            "project_id": project_id,
            "requester_id": user.id,
            "status": {"$in": ["QUEUED", "RUNNING"]},
            "cancel_requested": {"$ne": True},
        },
        {
            "$set": {
                "status": "RUNNING",
                "worker_job_id": worker_job_id,
                "started_at": datetime.now(timezone.utc),
                "updated_at": datetime.now(timezone.utc),
            }
        },
    )
    if not claimed.matched_count:
        latest = await database.value.document_summary_jobs.find_one({"_id": job_id})
        if latest and latest.get("status") in {"COMPLETED", "INCOMPLETE", "FAILED", "CANCELLED"}:
            return {"status": "COMPLETED", "summary_job_id": job_id, "summary_status": latest["status"]}
        raise RuntimeError("DOCUMENT_SUMMARY_JOB_CLAIM_FAILED")
    from src.schemas.contracts.requirements import RequirementDocumentSummaryRequest

    request = RequirementDocumentSummaryRequest.model_validate(job["request"])
    try:
        result = await summarize_project_documents(project_id, request, user, job_id)
        status = result["status"]
        await _update_summary_job(
            job_id,
            {
                "status": status,
                "result": result,
                "completed_at": datetime.now(timezone.utc),
                "progress": {
                    "processed": len(result.get("results", [])) + len(result.get("excluded", [])),
                    "total": len(result.get("results", [])) + len(result.get("excluded", [])),
                },
            },
        )
    except SummaryJobCancelled:
        status = "CANCELLED"
        await _update_summary_job(job_id, {"status": status, "completed_at": datetime.now(timezone.utc)})
    except Exception as error:
        status = "FAILED"
        await _update_summary_job(
            job_id,
            {
                "status": status,
                "error_code": type(error).__name__,
                "completed_at": datetime.now(timezone.utc),
            },
        )
    return {"status": "COMPLETED", "summary_job_id": job_id, "summary_status": status}


async def start_document_summary_job(project_id, request, user):
    await get_project(project_id, user, "requirement_document.extract")
    collection = database.value.document_summary_jobs
    existing = await collection.find_one(
        {
            "project_id": project_id,
            "requester_id": user.id,
            "idempotency_key": request.idempotency_key,
        }
    )
    if existing:
        return existing
    job_id = str(uuid4())
    now = datetime.now(timezone.utc)
    job = {
        "_id": job_id,
        "project_id": project_id,
        "requester_id": user.id,
        "idempotency_key": request.idempotency_key,
        "request": request.model_dump(),
        "locked_sources": await _lock_summary_sources(project_id, request, user),
        "status": "QUEUED",
        "progress": {"processed": 0, "total": 0},
        "cancel_requested": False,
        "created_at": now,
        "updated_at": now,
    }
    try:
        await collection.insert_one(job)
    except DuplicateKeyError:
        existing = await collection.find_one(
            {
                "project_id": project_id,
                "requester_id": user.id,
                "idempotency_key": request.idempotency_key,
            }
        )
        if existing:
            return existing
        raise
    try:
        await _enqueue_summary_job(job, user)
    except WorkerClientError:
        await collection.update_one(
            {"_id": job_id},
            {
                "$set": {
                    "status": "FAILED",
                    "error_code": "SUMMARY_WORKER_UNAVAILABLE",
                    "completed_at": datetime.now(timezone.utc),
                    "updated_at": datetime.now(timezone.utc),
                }
            },
        )
        return await collection.find_one({"_id": job_id})
    return await collection.find_one({"_id": job_id})


async def get_document_summary_job(project_id, job_id, user):
    await get_project(project_id, user, "requirement_document.read")
    job = await database.value.document_summary_jobs.find_one(
        {"_id": job_id, "project_id": project_id, "requester_id": user.id}
    )
    if not job:
        raise HTTPException(status_code=404, detail={"code": "DOCUMENT_SUMMARY_JOB_NOT_FOUND"})
    if job.get("status") in {"QUEUED", "RUNNING"} and job.get("worker_job_id"):
        try:
            worker_job = await worker_client.get(job["worker_job_id"])
            worker_status = str(worker_job.get("status") or "").lower()
            if worker_status == "failed":
                await _update_summary_job(
                    job_id,
                    {
                        "status": "FAILED",
                        "error_code": worker_job.get("error_code") or "SUMMARY_WORKER_FAILED",
                        "completed_at": datetime.now(timezone.utc),
                    },
                )
                job["status"] = "FAILED"
                job["error_code"] = worker_job.get("error_code") or "SUMMARY_WORKER_FAILED"
            elif worker_status == "canceled":
                await _update_summary_job(
                    job_id,
                    {"status": "CANCELLED", "completed_at": datetime.now(timezone.utc)},
                )
                job["status"] = "CANCELLED"
        except WorkerClientError:
            pass
    result = job.get("result") or {}
    for source in result.get("results", []):
        document = await get_requirement_document_record(source["document_id"], user)
        if (
            document.get("content_hash") != source.get("content_hash")
            or document.get("revision") != source.get("revision")
            or (document.get("source_version") or "1") != source.get("source_version")
            or document.get("ai_enabled") is not True
        ):
            raise HTTPException(status_code=409, detail={"code": "SUMMARY_SOURCE_CHANGED_OR_ACCESS_REVOKED"})
    return job


async def cancel_document_summary_job(project_id, job_id, user):
    await get_project(project_id, user, "requirement_document.extract")
    job_status = await database.value.document_summary_jobs.find_one(
        {"_id": job_id, "project_id": project_id, "requester_id": user.id},
        {"status": 1},
    )
    if not job_status or job_status.get("status") not in {"QUEUED", "RUNNING"}:
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_SUMMARY_JOB_NOT_CANCELLABLE"})
    result = await database.value.document_summary_jobs.update_one(
        {
            "_id": job_id,
            "project_id": project_id,
            "requester_id": user.id,
            "status": {"$in": ["QUEUED", "RUNNING"]},
        },
        {
            "$set": {
                "cancel_requested": True,
                "status": "CANCELLED" if job_status.get("status") == "QUEUED" else "RUNNING",
                "updated_at": datetime.now(timezone.utc),
            }
        },
    )
    if not result.matched_count:
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_SUMMARY_JOB_NOT_CANCELLABLE"})
    if job_status.get("status") == "QUEUED":
        job = await database.value.document_summary_jobs.find_one({"_id": job_id})
        try:
            if job and job.get("worker_job_id"):
                await worker_client.cancel(job["worker_job_id"])
        except WorkerClientError:
            pass
    return {"job_id": job_id, "status": "CANCEL_REQUESTED"}


async def retry_document_summary_job(project_id, job_id, user):
    await get_project(project_id, user, "requirement_document.extract")
    collection = database.value.document_summary_jobs
    job = await collection.find_one(
        {
            "_id": job_id,
            "project_id": project_id,
            "requester_id": user.id,
            "status": {"$in": ["FAILED", "INCOMPLETE"]},
        }
    )
    if not job:
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_SUMMARY_JOB_NOT_RETRYABLE"})
    for source in job.get("locked_sources") or []:
        try:
            document = await get_requirement_document_record(source["document_id"], user)
        except HTTPException as error:
            raise HTTPException(
                status_code=409,
                detail={"code": "SUMMARY_SOURCE_CHANGED_OR_ACCESS_REVOKED"},
            ) from error
        if (
            document.get("revision") != source.get("revision")
            or document.get("content_hash") != source.get("content_hash")
            or (document.get("source_version") or "1") != source.get("source_version")
            or (document.get("ai_enabled") is True) != source.get("ai_enabled")
        ):
            raise HTTPException(status_code=409, detail={"code": "SUMMARY_SOURCE_CHANGED_OR_ACCESS_REVOKED"})
    generation = int(job.get("dispatch_generation", 0)) + 1
    updated = await collection.update_one(
        {"_id": job_id, "status": job["status"]},
        {
            "$set": {
                "status": "QUEUED",
                "cancel_requested": False,
                "dispatch_generation": generation,
                "progress": {"processed": 0, "total": 0},
                "updated_at": datetime.now(timezone.utc),
            },
            "$unset": {"result": "", "error_code": "", "completed_at": ""},
        },
    )
    if not updated.modified_count:
        raise HTTPException(status_code=409, detail={"code": "DOCUMENT_SUMMARY_JOB_RUNNING"})
    job.update({"dispatch_generation": generation, "status": "QUEUED"})
    try:
        await _enqueue_summary_job(job, user, generation)
    except WorkerClientError:
        await collection.update_one(
            {"_id": job_id, "status": "QUEUED"},
            {"$set": {"status": "FAILED", "error_code": "SUMMARY_WORKER_UNAVAILABLE", "updated_at": datetime.now(timezone.utc)}},
        )
        return {"job_id": job_id, "status": "FAILED", "error_code": "SUMMARY_WORKER_UNAVAILABLE"}
    return {"job_id": job_id, "status": "QUEUED"}


async def summarize_project_documents(project_id, request, user, job_id=None):
    await get_project(project_id, user, "requirement_document.extract")
    results = []
    excluded = []
    job = (
        await database.value.document_summary_jobs.find_one({"_id": job_id}, {"locked_sources": 1})
        if job_id
        else None
    )
    locked_sources = (job or {}).get("locked_sources")
    if locked_sources is not None:
        selected = []
        for source in locked_sources:
            try:
                document = await get_requirement_document_record(source["document_id"], user)
            except HTTPException:
                excluded.append({"document_id": source["document_id"], "reason": "SOURCE_UNAVAILABLE"})
                continue
            if (
                document.get("revision") != source.get("revision")
                or document.get("content_hash") != source.get("content_hash")
                or (document.get("source_version") or "1") != source.get("source_version")
                or (document.get("ai_enabled") is True) != source.get("ai_enabled")
            ):
                excluded.append({"document_id": source["document_id"], "reason": "SOURCE_CHANGED"})
                continue
            selected.append(document)
    elif request.all_readable:
        documents = await requirement_document_repository.list({"project_id": project_id}, None)
        documents = [document for document in documents if can_read_requirement_document(document, user)]
        selected = [document for document in documents if document.get("status") != "ARCHIVED"]
    else:
        selected = []
        for document_id in request.document_ids:
            document = await get_requirement_document_record(document_id, user)
            if str(document.get("project_id")) != str(project_id):
                raise HTTPException(status_code=404, detail={"code": "DOCUMENT_NOT_FOUND"})
            selected.append(document)
    total = len(selected)
    await _update_summary_job(job_id, {"progress": {"processed": 0, "total": total}})
    for document in selected:
        await _check_summary_job(job_id)
        if document.get("status") != "READY":
            excluded.append({"document_id": str(document["_id"]), "reason": "NOT_READY"})
            await _update_summary_job(job_id, {"progress": {"processed": len(results) + len(excluded), "total": total}})
            continue
        if document.get("ai_enabled") is not True:
            excluded.append({"document_id": str(document["_id"]), "reason": "AI_READ_NOT_ALLOWED"})
            await _update_summary_job(job_id, {"progress": {"processed": len(results) + len(excluded), "total": total}})
            continue
        try:
            content = await _document_text(document)
        except SummaryJobCancelled:
            raise
        except Exception as error:
            excluded.append(
                {
                    "document_id": str(document["_id"]),
                    "reason": "CONTENT_READ_FAILED",
                    "error": type(error).__name__,
                }
            )
            continue
        if not content:
            excluded.append({"document_id": str(document["_id"]), "reason": "CONTENT_UNAVAILABLE"})
            await _update_summary_job(job_id, {"progress": {"processed": len(results) + len(excluded), "total": total}})
            continue
        chunks = split_source_text(content)
        if not chunks:
            excluded.append({"document_id": str(document["_id"]), "reason": "CONTENT_UNAVAILABLE"})
            await _update_summary_job(job_id, {"progress": {"processed": len(results) + len(excluded), "total": total}})
            continue
        cache_key = hashlib.sha256(
            "|".join(
                [
                    project_id,
                    str(document["_id"]),
                    str(document.get("content_hash") or ""),
                    str(document.get("source_version") or "1"),
                    str(document.get("revision") or ""),
                    request.output_language,
                    SUMMARY_PROMPT_VERSION,
                ]
            ).encode("utf-8")
        ).hexdigest()
        try:
            cached = await database.value.document_summary_cache.find_one({"_id": cache_key})
            summary = cached.get("summary") if cached else None
            if not summary:
                summaries = []
                for chunk_index, chunk in enumerate(chunks):
                    await _check_summary_job(job_id)
                    await _update_summary_job(job_id, {"phase": "SUMMARIZING_DOCUMENT"})
                    await _assert_summary_source_current(document, user)
                    chunk_key = hashlib.sha256(
                        "|".join(
                            [
                                cache_key,
                                str(chunk_index),
                                hashlib.sha256(chunk.encode("utf-8")).hexdigest(),
                            ]
                        ).encode("utf-8")
                    ).hexdigest()
                    cached_chunk = await database.value.document_summary_chunk_cache.find_one(
                        {"_id": chunk_key}
                    )
                    if cached_chunk:
                        summaries.append(cached_chunk["summary"])
                        await _update_summary_job(
                            job_id,
                            {
                                "progress": {
                                    "processed": len(results) + len(excluded),
                                    "total": total,
                                    "current_document": document.get("filename"),
                                    "chunks_processed": chunk_index + 1,
                                    "chunks_total": len(chunks),
                                }
                            },
                        )
                        continue
                    chunk_summary = await ai_automation_client.summarize_knowledge_text(
                        chunk, request.output_language
                    )
                    await database.value.document_summary_chunk_cache.replace_one(
                        {"_id": chunk_key},
                        {
                            "_id": chunk_key,
                            "project_id": project_id,
                            "document_id": str(document["_id"]),
                            "source_hash": document.get("content_hash"),
                            "output_language": request.output_language,
                            "prompt_version": SUMMARY_PROMPT_VERSION,
                            "chunk_index": chunk_index,
                            "chunk_hash": hashlib.sha256(chunk.encode("utf-8")).hexdigest(),
                            "summary": chunk_summary,
                        },
                        upsert=True,
                    )
                    summaries.append(chunk_summary)
                    await _update_summary_job(
                        job_id,
                        {
                            "progress": {
                                "processed": len(results) + len(excluded),
                                "total": total,
                                "current_document": document.get("filename"),
                                "chunks_processed": chunk_index + 1,
                                "chunks_total": len(chunks),
                            }
                        },
                    )
                while len("\n\n".join(summaries)) > SUMMARY_INPUT_LIMIT:
                    combined = "\n\n".join(summaries)
                    summaries = [
                        await ai_automation_client.summarize_knowledge_text(
                            combined[index : index + SUMMARY_INPUT_LIMIT], request.output_language
                        )
                        for index in range(0, len(combined), SUMMARY_INPUT_LIMIT)
                    ]
                summary = (
                    summaries[0]
                    if len(summaries) == 1
                    else await ai_automation_client.summarize_knowledge_text(
                        "\n\n".join(summaries), request.output_language
                    )
                )
            await _assert_summary_source_current(document, user)
            if not cached:
                await database.value.document_summary_cache.replace_one(
                    {"_id": cache_key},
                    {
                        "_id": cache_key,
                        "project_id": project_id,
                        "document_id": str(document["_id"]),
                        "source_hash": document.get("content_hash"),
                        "output_language": request.output_language,
                        "prompt_version": SUMMARY_PROMPT_VERSION,
                        "summary": summary,
                    },
                    upsert=True,
                )
            report = document.get("parse_report") or {}
            results.append(
                {
                    "document_id": str(document["_id"]),
                    "filename": document.get("filename"),
                    "source_version": document.get("source_version") or "1",
                    "revision": document.get("revision"),
                    "content_hash": document.get("content_hash"),
                    "summary": summary,
                    "complete": not bool(
                        report.get("unprocessed_blocks") or report.get("extraction_warnings")
                    ),
                    "unprocessed_blocks": report.get("unprocessed_blocks") or [],
                    "page_count": report.get("page_count"),
                }
            )
            await _update_summary_job(
                job_id, {"progress": {"processed": len(results) + len(excluded), "total": total}}
            )
        except Exception as error:
            excluded.append(
                {"document_id": str(document["_id"]), "reason": "SUMMARY_FAILED", "error": type(error).__name__}
            )
            await _update_summary_job(
                job_id, {"progress": {"processed": len(results) + len(excluded), "total": total}}
            )
    if not results:
        return {
            "status": "FAILED",
            "results": [],
            "excluded": excluded,
            "summary": "",
        }
    summaries = [f"{item['filename']} (v{item['source_version']}): {item['summary']}" for item in results]
    combined = "\n\n".join(summaries)
    while len(combined) > SUMMARY_INPUT_LIMIT:
        await _check_summary_job(job_id)
        combined = "\n\n".join(
            [
                await ai_automation_client.summarize_knowledge_text(
                    combined[index : index + SUMMARY_INPUT_LIMIT], request.output_language
                )
                for index in range(0, len(combined), SUMMARY_INPUT_LIMIT)
            ]
        )
    for item in results:
        await _check_summary_job(job_id)
        source = await get_requirement_document_record(item["document_id"], user)
        if (
            source.get("content_hash") != item.get("content_hash")
            or source.get("revision") != item.get("revision")
            or (source.get("source_version") or "1") != item.get("source_version")
            or source.get("ai_enabled") is not True
        ):
            raise HTTPException(status_code=409, detail={"code": "SUMMARY_SOURCE_CHANGED_OR_ACCESS_REVOKED"})
    overall = (
        combined
        if len(results) == 1
        else await ai_automation_client.summarize_knowledge_text(combined, request.output_language)
    )
    return {
        "status": "COMPLETED" if not excluded and all(item["complete"] for item in results) else "INCOMPLETE",
        "results": results,
        "excluded": excluded,
        "summary": overall,
    }
